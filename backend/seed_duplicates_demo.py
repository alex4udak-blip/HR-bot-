"""Наполнение ЛОКАЛЬНОЙ базы парами кандидатов-дублей — чтобы было что проверять.

Создаёт по паре на КАЖДЫЙ тир единого матчера (services/duplicate_matcher):
точная почта, телефон в разном формате, ФИО с отчеством и транслитом, та же
ссылка на резюме, личный telegram, мягкий тир (разное ФИО, но те же 7 цифр
телефона и дата рождения), совпавший текст резюме и архивный дубль.

У «старой» анкеты каждой пары заполнены статус/причина отказа/история этапов и
резюме — чтобы окно сравнения было не пустым. Помечены маркером
``extra_data.demo_dup_fixture`` и удаляются одной командой:

    cd backend && set -a && . ./dev.env && set +a && .venv/bin/python seed_duplicates_demo.py
    cd backend && set -a && . ./dev.env && set +a && .venv/bin/python seed_duplicates_demo.py --purge

Идемпотентен: повторный запуск сначала сносит прошлую партию. Не для прода —
отказывается работать, если DATABASE_URL не смотрит на localhost.
"""
import asyncio
import os
import sys
from datetime import datetime, timedelta

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession, create_async_engine, async_sessionmaker

from api.models.database import (
    Entity, EntityFile, EntityFileType, EntityType, EntityStatus, Organization, User,
)
from api.services.similarity import detect_archived_duplicate, transliterate_ru_to_en

DB_URL = os.environ.get("DATABASE_URL", "")
MARKER = "demo_dup_fixture"

_TODAY = datetime.now()


def _d(days_ago: int) -> str:
    return (_TODAY - timedelta(days=days_ago)).strftime("%d.%m.%Y")


# Резюме и история — только у «старой» анкеты: так в окне сравнения видно, ЧТО
# теряется при неверном решении «это разные люди».
def _old_extra(**kw) -> dict:
    base = {
        MARKER: True,
        "timeline_events": [
            {"date": _d(120), "title": "Добавлен в базу"},
            {"date": _d(100), "title": "Скрининг"},
            {"date": _d(80), "title": "Техническое интервью"},
        ],
        "notes": [
            {"text": "Созвон был, просил вернуться через полгода.",
             "author_name": "Настя", "date": (_TODAY - timedelta(days=80)).isoformat()},
        ],
    }
    base.update(kw)
    return base


def _new_extra(**kw) -> dict:
    base = {MARKER: True}
    base.update(kw)
    return base


# (пояснение, «старая» анкета, «новая» анкета) — новая создаётся ПОСЛЕ старой,
# на ней и сработает детект.
PAIRS = [
    (
        "Точная почта (тир email, 100%)",
        dict(name="Гончарова Алиса", position="UX-дизайнер", company="Ozon",
             email="alisa.goncharova@gmail.com", phone="+7 903 118-44-01",
             status=EntityStatus.rejected,
             extra_data=_old_extra(city="Москва", birth_date="1994-03-12",
                                   rejection_reason="не сошлись по вилке")),
        dict(name="Алиса Гончарова", position="Product Designer", company="Wildberries",
             email="alisa.goncharova@gmail.com", phone="+7 903 118-44-01",
             extra_data=_new_extra(city="Москва", birth_date="1994-03-12")),
    ),
    (
        "Почта сменила домен, локаль та же (тир email)",
        dict(name="Тарасов Вадим", position="Backend-разработчик", company="Тинькофф",
             email="v.tarasov@gmail.com",
             extra_data=_old_extra(city="Казань")),
        dict(name="Вадим Тарасов", position="Python Developer", company="Яндекс",
             email="v.tarasov@mail.ru",
             extra_data=_new_extra(city="Казань")),
    ),
    (
        "Телефон в другом формате (тир phone)",
        dict(name="Ильина Марина", position="Рекрутёр", company="Сбер",
             phone="8 (912) 345-67-89",
             extra_data=_old_extra(city="Екатеринбург")),
        # Фамилия сменилась (замужество) — по ФИО пара не матчится, ловит только телефон.
        dict(name="Соколова Марина", position="Talent Acquisition", company="СберТех",
             phone="+7 912 3456789",
             extra_data=_new_extra(city="Екатеринбург")),
    ),
    (
        "ФИО с отчеством + транслит, контактов нет (тир name)",
        dict(name="Векленко Кирилл Дмитриевич", position="Аналитик", company="X5",
             extra_data=_old_extra(city="Москва")),
        dict(name="Kirill Veklenko", position="Data Analyst", company="Lamoda",
             extra_data=_new_extra(city="Москва")),
    ),
    (
        "Личный telegram (тир telegram)",
        dict(name="Шарипов Тимур", position="QA-инженер", company="МТС",
             telegram_usernames=["timur_qa_2024"],
             extra_data=_old_extra(city="Уфа", )),
        dict(name="Тимур Ш.", position="QA Automation", company="VK",
             telegram_usernames=["@Timur_QA_2024"],
             extra_data=_new_extra(city="Уфа")),
    ),
    (
        "Та же ссылка на резюме hh (тир source)",
        dict(name="Кандидат без имени", position="Flutter Developer, Минск, 27 лет",
             extra_data=_old_extra(source_url="https://hh.ru/resume/a1b2c3d4e5f6?hhtmFrom=chat&t=111")),
        dict(name="Flutter Developer, Минск", position="Flutter Developer",
             extra_data=_new_extra(source_url="https://hh.ru/resume/a1b2c3d4e5f6?vacancyId=99&t=999")),
    ),
    (
        "Мягкий тир: другое ФИО, те же 7 цифр телефона + дата рождения",
        dict(name="Петров Александр", position="Маркетолог", company="Avito",
             phone="+7 916 000-11-22",
             extra_data=_old_extra(birth_date="1990-05-14", city="Москва")),
        dict(name="Сидоров Иван", position="Head of Growth", company="Ozon",
             phone="+7 495 000-11-22",
             extra_data=_new_extra(birth_date="14.05.1990", city="Москва")),
    ),
    (
        "Мягкий тир 75%: те же 7 цифр телефона + дата рождения, ГОРОДА РАЗНЫЕ",
        dict(name="Романова Ольга", position="Логист", company="СДЭК",
             phone="+7 921 777-31-40",
             extra_data=_old_extra(birth_date="1988-11-02", city="Санкт-Петербург")),
        dict(name="Кузнецова Алина", position="Менеджер ВЭД", company="Байкал-Сервис",
             phone="+7 495 777-31-40",
             extra_data=_new_extra(birth_date="02.11.1988", city="Москва")),
    ),
    (
        "НИЖЕ ПОРОГА (48 баллов): та же дата рождения + тот же город — не показывается вовсе",
        dict(name="Зайцев Артём", position="Копирайтер", company="Skyeng",
             phone="+7 911 100-20-30",
             extra_data=_old_extra(birth_date="1995-07-19", city="Казань")),
        dict(name="Морозов Денис", position="Редактор", company="Нетология",
             phone="+7 999 400-50-60",
             extra_data=_new_extra(birth_date="19.07.1995", city="Казань")),
    ),
    (
        "Текст резюме совпал ЧАСТИЧНО (инфо-тир text — слияние НЕ предлагается)",
        dict(name="Никитин Егор", position="SMM-менеджер", company="Самокат",
             email="egor.nikitin@yandex.ru",
             extra_data=_old_extra(
                 city="Санкт-Петербург",
                 summary="Запускал и вёл рекламные кампании в Facebook Ads и VK Ads, отвечал "
                         "за медиаплан, работал с блогерами и подрядчиками, считал юнит-экономику "
                         "и еженедельно отчитывался по ROMI перед руководителем отдела маркетинга.",
             )),
        dict(name="Соловьёв Никита", position="SMM", company="Додо Пицца",
             email="nikita.solovyov@yandex.ru",
             extra_data=_new_extra(
                 city="Москва",
                 # Тот же текст, но без последней фразы — совпадение 86%, а не 100%:
                 # в интерфейсе видно промежуточный процент текстового тира.
                 summary="Запускал и вёл рекламные кампании в Facebook Ads и VK Ads, отвечал "
                         "за медиаплан, работал с блогерами и подрядчиками, считал юнит-экономику "
                         "и еженедельно отчитывался по ROMI.",
             )),
    ),
    (
        "Дубль лежит в АРХИВЕ (виден только с include_archived)",
        dict(name="Жукова Ольга", position="Бухгалтер", company="Лента",
             email="olga.zhukova@bk.ru", phone="+7 921 555-30-10", is_archived=True,
             extra_data=_old_extra(city="Санкт-Петербург",
                                   description="Импорт ClickUp: бухгалтерия, 1С, 9 лет.")),
        dict(name="Ольга Жукова", position="Главный бухгалтер", company="Пятёрочка",
             email="olga.zhukova@bk.ru",
             extra_data=_new_extra(city="Санкт-Петербург")),
    ),
]


# Кейс Эльвиры (2026-09-16): «Иванов Кирилл Владимирович» поднимал ЧЕТЫРЁХ
# однофамильцев-тёзок с разными отчествами, и все шли как «точное совпадение».
# Проверяем, что теперь предлагаются только настоящие дубли (совпало отчество и
# совпал телефон), а трое с чужими отчествами не предлагаются вовсе.
NAMESAKES = [
    # (анкета, ожидание)
    (dict(name="Иванов Кирилл Владимирович", position="Логист", company="Деловые линии",
          phone="+7 917 200-40-60",
          extra_data=_old_extra(city="Нижний Новгород")),
     "дубль: то же отчество"),
    (dict(name="Кирилл Иванов", position="Специалист по логистике", company="СДЭК",
          phone="8 917 2004060",
          extra_data=_old_extra(city="Нижний Новгород")),
     "дубль: тот же телефон"),
    (dict(name="Иванов Кирилл Евгеньевич", position="Водитель", company="Магнит",
          extra_data=_old_extra(city="Краснодар")), "однофамилец: другое отчество"),
    (dict(name="Иванов Кирилл Сергеевич", position="Менеджер", company="Ozon",
          extra_data=_old_extra(city="Москва")), "однофамилец: другое отчество"),
    (dict(name="Иванов Кирилл Петрович", position="Сварщик", company="ММК",
          extra_data=_old_extra(city="Магнитогорск")), "однофамилец: другое отчество"),
]
NAMESAKE_NEW = dict(name="Иванов Кирилл Владимирович", position="Руководитель склада",
                    company="Wildberries", phone="+7 917 200-40-60",
                    extra_data=_new_extra(city="Нижний Новгород"))


# Шрифт с кириллицей: base-14 Helvetica её не кодирует, и резюме получалось
# латиницей («cv: goncharova alisa») — сравнивать такое глазами невозможно.
# Порядок важен: Arial Unicode весит 22 МБ и вшивается в каждый файл, поэтому он
# последний — сначала обычные шрифты с кириллицей.
_FONT_CANDIDATES = (
    "/System/Library/Fonts/Supplemental/Arial.ttf",
    "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    "/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf",
    "/System/Library/Fonts/Supplemental/Arial Unicode.ttf",
    "/Library/Fonts/Arial Unicode.ttf",
)


def _font_path():
    for path in _FONT_CANDIDATES:
        if os.path.exists(path):
            return path
    return None


# Блоки резюме: (вид, текст). Вид задаёт кегль и отступ при отрисовке.
TITLE, SUB, HEADING, BODY, BULLET = "title", "sub", "heading", "body", "bullet"
_STYLE = {
    TITLE: (18, 0, 26),
    SUB: (11, 0, 18),
    HEADING: (12, 0, 24),
    BODY: (10, 0, 15),
    BULLET: (10, 14, 15),
}


def _render_resume_pdf(blocks) -> bytes:
    """Резюме в PDF с текстовым слоем: A4, разделы, переносы, вторая страница при
    необходимости. Текст извлекается pdfplumber'ом — именно он идёт в сравнение."""
    import pymupdf

    font_file = _font_path()
    if font_file is None:  # деградация: латиница base-14, но файл валидный
        blocks = [(k, transliterate_ru_to_en(t)) for k, t in blocks]
    font = pymupdf.Font(fontfile=font_file) if font_file else pymupdf.Font("helv")
    fontname = "CV" if font_file else "helv"

    doc = pymupdf.open()
    page = doc.new_page()
    if font_file:
        page.insert_font(fontname=fontname, fontfile=font_file)
    left, right, top, bottom = 56, 539, 64, 780
    y = top

    def _wrap(text: str, size: float, width: float):
        words, line, out = text.split(), "", []
        for w in words:
            probe = f"{line} {w}".strip()
            if font.text_length(probe, size) <= width:
                line = probe
            else:
                if line:
                    out.append(line)
                line = w
        if line:
            out.append(line)
        return out or [""]

    for kind, text in blocks:
        size, indent, leading = _STYLE[kind]
        prefix = "• " if kind == BULLET else ""
        for i, line in enumerate(_wrap(prefix + text, size, right - left - indent)):
            if y > bottom:
                page = doc.new_page()
                if font_file:
                    page.insert_font(fontname=fontname, fontfile=font_file)
                y = top
            page.insert_text((left + indent + (10 if i and kind == BULLET else 0), y),
                             line, fontsize=size, fontname=fontname)
            y += leading
        y += 4 if kind in (TITLE, HEADING) else 0
    # Без subset шрифт вшивается целиком: демо-резюме весило 23 МБ на файл.
    try:
        doc.subset_fonts()
    except Exception:
        pass
    return doc.tobytes()


# Наполнение резюме. Каждый кандидат получает СВОЙ текст (обязанности,
# прошлое место, вуз и навыки выбираются по хэшу имени) — иначе детектор
# копипаста находил бы «близнецов» у всех демо-анкет сразу.
_PREV_COMPANIES = ["Яндекс", "СберТех", "Ozon", "Авито", "Тинькофф", "Lamoda", "Самокат", "VK"]
_UNIVERSITIES = [
    "МГУ им. Ломоносова, экономический факультет",
    "НИУ ВШЭ, факультет бизнеса и менеджмента",
    "СПбГУ, факультет психологии",
    "РАНХиГС, управление персоналом",
    "КФУ, институт управления и экономики",
]
_DUTIES = [
    "Вёл полный цикл подбора: от снятия заявки до выхода кандидата.",
    "Выстроил воронку найма с нуля и еженедельно считал конверсию этапов.",
    "Проводил скрининги и финальные интервью вместе с нанимающими менеджерами.",
    "Согласовывал офферы с финансами, закрывал вилки вне утверждённой сетки.",
    "Перевёл отчётность на дашборды, сократил ручную работу команды вдвое.",
    "Запустил реферальную программу, доля рекомендаций выросла до трети найма.",
    "Вёл базу кандидатов: чистил дубли, следил за актуальностью контактов.",
    "Договаривался с подрядчиками и снизил стоимость привлечения на 20%.",
]
_ACHIEVEMENTS = [
    "Сократил время закрытия вакансии с 54 до 31 дня.",
    "Закрыл 42 вакансии за год при плане 30.",
    "Поднял долю принятых офферов с 68% до 85%.",
    "Собрал команду из 12 человек под запуск нового направления.",
]
_ABOUT = [
    "Работаю в найме девятый год, последние четыре — в {town}. Люблю прозрачные "
    "процессы: считаю метрики, не боюсь говорить нанимающим менеджерам «нет» и "
    "аккуратно веду базу — половина проблем найма начинается с грязных данных.",
    "Пришла в подбор из операционки, поэтому смотрю на найм как на процесс с "
    "узкими местами. В {comp} собрала отчётность, по которой стало видно, где "
    "воронка теряет людей, и починила два самых дорогих этапа.",
    "Больше всего люблю сложный точечный поиск: когда кандидатов на рынке "
    "двадцать человек и до каждого нужно достучаться лично. Умею писать письма, "
    "на которые отвечают, и не выгорать от отказов.",
    "Считаю, что рекрутёр отвечает не за количество собеседований, а за то, "
    "чтобы человек вышел и остался. Поэтому довожу кандидата до конца "
    "испытательного и собираю обратную связь с обеих сторон.",
]
_COURSES = [
    "«Оценка персонала» (2020)", "«Аналитика найма» (2022)",
    "«Интервью по компетенциям» (2019)", "«HR-аналитика на SQL» (2023)",
    "«Employer brand» (2021)", "«Executive search» (2018)",
]
_FORMATS = [
    "Английский — Upper-Intermediate. Готов к гибридному формату.",
    "Английский — B1, читаю профильную литературу. Рассматриваю удалёнку.",
    "Английский — Advanced, вёл найм в международной команде. Офис или гибрид.",
]
_SKILLS = [
    "поиск и подбор, executive search, массовый подбор",
    "оценка компетенций, структурированное интервью, кейс-интервью",
    "ATS, аналитика воронки, отчётность по метрикам найма",
    "работа с нанимающими менеджерами, калибровка требований",
    "HR-бренд, работа с площадками и реферальной программой",
]


def _resume_blocks(name, position, company, city, phone=None, email=None, birth=None) -> list:
    """Полноценное резюме: шапка, опыт с двумя местами работы, образование,
    навыки и «о себе». ~2000 знаков — на таком тексте сравнение уже осмысленно."""
    import hashlib

    import random

    seed = int(hashlib.md5((name or "").encode()).hexdigest(), 16)
    rnd = random.Random(seed)
    pos = position or "Специалист"
    comp = company or "—"
    town = city or "Москва"
    # Независимые выборки (не арифметика по одному числу): при выборе «по модулю»
    # разные люди слишком часто получали один и тот же набор фраз, и детектор
    # копипаста честно ловил их как близнецов — проверять было не на чем.
    prev = rnd.choice(_PREV_COMPANIES)
    uni = rnd.choice(_UNIVERSITIES)
    duties = rnd.sample(_DUTIES, 3)
    duties_prev = rnd.sample([d for d in _DUTIES if d not in duties], 2)
    ach = rnd.choice(_ACHIEVEMENTS)
    skills = ", ".join(rnd.sample(_SKILLS, 3))
    about = rnd.choice(_ABOUT).format(town=town, comp=comp)
    courses = ", ".join(rnd.sample(_COURSES, 2))
    fmt = rnd.choice(_FORMATS)
    start_year = rnd.choice((2019, 2020, 2021, 2022))
    prev_from, prev_to = start_year - 4, start_year

    contacts = " · ".join(x for x in (phone, email, town, f"д.р. {birth}" if birth else None) if x)
    return [
        (TITLE, name),
        (SUB, f"{pos} · {comp}"),
        (SUB, contacts or town),
        (HEADING, "ОПЫТ РАБОТЫ"),
        (BODY, f"{start_year} — настоящее время · {comp} · {pos}"),
        *[(BULLET, d) for d in duties],
        (BULLET, ach),
        (BODY, f"{prev_from} — {prev_to} · {prev} · Специалист по подбору персонала"),
        *[(BULLET, d) for d in duties_prev],
        (HEADING, "ОБРАЗОВАНИЕ"),
        (BODY, f"{prev_from - 5} — {prev_from - 1} · {uni}"),
        (BODY, f"Курсы: {courses}"),
        (HEADING, "КЛЮЧЕВЫЕ НАВЫКИ"),
        (BODY, skills),
        (HEADING, "О СЕБЕ"),
        (BODY, about),
        (BODY, fmt),
    ]


def _scan_pdf(blocks) -> bytes:
    """PDF-СКАН: страница отрендерена в картинку, текстового слоя нет."""
    import pymupdf

    src = pymupdf.open(stream=_render_resume_pdf(blocks), filetype="pdf")
    pix = src[0].get_pixmap(dpi=130)
    out = pymupdf.open()
    page = out.new_page(width=pix.width, height=pix.height)
    page.insert_image(page.rect, stream=pix.tobytes("png"))
    return out.tobytes()


def _docx_bytes(blocks) -> bytes:
    """То же резюме в DOCX — второй формат, который читает парсер."""
    import io

    from docx import Document

    doc = Document()
    for kind, text in blocks:
        if kind == TITLE:
            doc.add_heading(text, level=1)
        elif kind == HEADING:
            doc.add_heading(text, level=2)
        elif kind == BULLET:
            doc.add_paragraph(text, style="List Bullet")
        else:
            doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _attach_file(db: AsyncSession, entity: Entity, org_id: int, data: bytes,
                 file_name: str, mime: str) -> None:
    db.add(EntityFile(
        entity_id=entity.id, org_id=org_id, file_type=EntityFileType.resume,
        file_name=file_name, file_data=data, file_size=len(data),
        mime_type=mime, description="Демо-резюме для проверки извлечения текста",
    ))


def _attach_resume(db: AsyncSession, entity: Entity, org_id: int, blocks) -> None:
    """Прикрепить кандидату PDF-резюме (file_type=resume) — как загрузка файла в
    карточке: содержимое лежит в БД (file_data)."""
    _attach_file(db, entity, org_id, _render_resume_pdf(blocks),
                 f"resume_{entity.id}.pdf", "application/pdf")


# ---------------------------------------------------------------------------
# Блок «проверка текста резюме» (владелец 17.09.2026): файлы прикладываем, а
# текст НЕ извлекаем — чтобы было на чём прогнать бэкфилл и увидеть разницу.
# ---------------------------------------------------------------------------

# (анкета, как приложить файл, что должно получиться)
RESUME_TEXT_CASES = [
    (dict(name="Ефимов Глеб", position="Senior Recruiter", company="Тинькофф",
          email="gleb.efimov@x.com", phone="+7 905 400-10-20",
          extra_data=_new_extra(city="Москва")),
     "pdf-shared", "PDF с текстовым слоем, текст ОБЩИЙ со следующим кандидатом"),
    (dict(name="Латыпова Динара", position="Talent Partner", company="Ozon",
          email="dinara.latypova@y.com", phone="+7 927 800-70-60",
          extra_data=_new_extra(city="Казань")),
     "pdf-shared", "тот же текст резюме, общих контактов НЕТ — ловится только текстом"),
    (dict(name="Сканов Пётр", position="Аналитик", company="X5",
          email="p.skanov@z.com",
          extra_data=_new_extra(city="Самара")),
     "scan", "скан-PDF: текстового слоя нет, OCR к PDF не подключён"),
    (dict(name="Докунова Дарья", position="HR BP", company="Авито",
          email="d.dokunova@z.com",
          extra_data=_new_extra(city="Новосибирск")),
     "docx", "резюме в DOCX — второй поддерживаемый формат"),
    (dict(name="Пустой Кандидат", position="Без резюме", company="—",
          email="empty.candidate@z.com",
          extra_data=_new_extra(city="Москва")),
     None, "файлов нет — СЮДА загрузите своё PDF и проверьте путь загрузки"),
]

async def _purge(db: AsyncSession, org_id: int) -> int:
    rows = (await db.execute(
        select(Entity).where(Entity.org_id == org_id, Entity.type == EntityType.candidate)
    )).scalars().all()
    killed = 0
    for e in rows:
        if isinstance(e.extra_data, dict) and e.extra_data.get(MARKER):
            await db.delete(e)
            killed += 1
    # Чужие карточки могли получить ссылку на удалённые — снимаем висячие флаги.
    ids = {e.id for e in rows if isinstance(e.extra_data, dict) and e.extra_data.get(MARKER)}
    for e in rows:
        if not isinstance(e.extra_data, dict) or e.extra_data.get(MARKER):
            continue
        if e.extra_data.get("hidden_duplicate_id") in ids:
            ne = dict(e.extra_data)
            ne.pop("hidden_duplicate_id", None)
            ne.pop("hidden_duplicate_meta", None)
            e.extra_data = ne
    await db.commit()
    return killed



async def _print_status(db: AsyncSession, org_id: int) -> None:
    """Что сейчас с текстом резюме у кандидатов организации: есть ли файл, достали
    ли из него текст, нашёлся ли близнец по тексту. Команда для повторных проверок:
    ``seed_duplicates_demo.py --status``."""
    from api.services.resume_text_twin import resume_text_blob, text_shingles

    ents = (await db.execute(
        select(Entity).where(Entity.org_id == org_id, Entity.type == EntityType.candidate)
        .order_by(Entity.id)
    )).scalars().all()
    files = (await db.execute(
        select(EntityFile.entity_id, EntityFile.file_name)
        .where(EntityFile.file_type == EntityFileType.resume)
    )).all()
    by_entity: dict = {}
    for eid, fname in files:
        by_entity.setdefault(eid, []).append(fname)

    print(f"{'id':>4}  {'кандидат':24} {'файл резюме':26} {'текст':>7}  {'сравним':7} близнец")
    print("-" * 92)
    for e in ents:
        names = by_entity.get(e.id) or []
        if not names and not (e.extra_data or {}).get("resume_text"):
            continue
        ed = e.extra_data if isinstance(e.extra_data, dict) else {}
        src = ed.get("resume_text_source") or {}
        chars = len(ed.get("resume_text") or "")
        comparable = "да" if len(text_shingles(resume_text_blob(ed))) >= 5 else "НЕТ"
        twin = (ed.get("text_twin") or {}).get("twin_id")
        mark = " (скан)" if src.get("empty") else ""
        print(f"{e.id:>4}  {e.name[:24]:24} {(names[0] if names else '—')[:26]:26} "
              f"{chars:>7}  {comparable:7} {twin or '—'}{mark}")


async def main() -> None:
    if "localhost" not in DB_URL and "127.0.0.1" not in DB_URL:
        raise SystemExit(f"Отказ: DATABASE_URL не локальный ({DB_URL!r}). Скрипт только для локальной базы.")

    engine = create_async_engine(DB_URL, echo=False)
    Session = async_sessionmaker(engine, class_=AsyncSession, expire_on_commit=False)

    async with Session() as db:
        org = (await db.execute(
            select(Organization).where(Organization.slug == "default")
        )).scalar_one_or_none()
        if org is None:
            raise SystemExit("Нет организации 'default'. Запусти бэкенд — он создаёт её на старте.")

        if "--status" in sys.argv:
            await _print_status(db, org.id)
            await engine.dispose()
            return

        killed = await _purge(db, org.id)
        if killed:
            print(f"Удалено прошлых демо-дублей: {killed}")
        if "--purge" in sys.argv:
            await engine.dispose()
            return

        author = (await db.execute(
            select(User).where(User.email == "recruiter.test@example.com")
        )).scalar_one_or_none()
        author_id = author.id if author else None

        print("Создаю пары дублей:")
        for title, old_kw, new_kw in PAIRS:
            old = Entity(org_id=org.id, type=EntityType.candidate, created_by=author_id,
                         status=old_kw.pop("status", EntityStatus.new), **old_kw)
            db.add(old)
            await db.flush()

            _attach_resume(db, old, org.id, _resume_blocks(
                old.name, old_kw.get("position"), old_kw.get("company"),
                (old_kw.get("extra_data") or {}).get("city"),
                phone=old_kw.get("phone"), email=old_kw.get("email"),
                birth=(old_kw.get("extra_data") or {}).get("birth_date"),
            ))

            new = Entity(org_id=org.id, type=EntityType.candidate, created_by=author_id,
                         status=EntityStatus.new, **new_kw)
            db.add(new)
            await db.flush()

            # Тот же путь, что у ручного добавления кандидата: детект личности +
            # детектор копипаста текста резюме.
            from api.services.resume_text_twin import detect_resume_text_twin
            match_id = await detect_archived_duplicate(db, new)
            if match_id:
                # Ровно как на пути создания кандидата (routes/entities/crud.py):
                # сам детект пишет только мету, а ссылку на пару ставит вызывающий.
                extra = dict(new.extra_data or {})
                extra["hidden_duplicate_id"] = match_id
                new.extra_data = extra
                await db.flush()
            twin_id, sim = await detect_resume_text_twin(db, new)
            meta = (new.extra_data or {}).get("hidden_duplicate_meta") or {}
            mark = (
                f"тир {meta.get('strength')} · {meta.get('confidence')}%"
                if match_id else (f"только текст ({round(sim * 100)}%)" if twin_id else "НЕ найден")
            )
            print(f"  [{old.id} ← {new.id}] {title}\n        {new.name}: {mark}")

        print("\nКейс Эльвиры — один новый кандидат против пяти однофамильцев:")
        for kw, note in NAMESAKES:
            old = Entity(org_id=org.id, type=EntityType.candidate, created_by=author_id,
                         status=EntityStatus.new, **kw)
            db.add(old)
            await db.flush()
            _attach_resume(db, old, org.id, _resume_blocks(
                old.name, kw.get("position"), kw.get("company"),
                (kw.get("extra_data") or {}).get("city"),
                phone=kw.get("phone"), email=kw.get("email"),
            ))
            print(f"  [{old.id}] {old.name:32} — {note}")
        new = Entity(org_id=org.id, type=EntityType.candidate, created_by=author_id,
                     status=EntityStatus.new, **NAMESAKE_NEW)
        db.add(new)
        await db.flush()
        match_id = await detect_archived_duplicate(db, new)
        if match_id:
            extra = dict(new.extra_data or {})
            extra["hidden_duplicate_id"] = match_id
            new.extra_data = extra
            await db.flush()
        from api.services.similarity import similarity_service
        offered = await similarity_service.detect_duplicates(db=db, entity=new)
        print(f"  [{new.id}] {new.name:32} — НОВЫЙ, предложено дублей: {len(offered)}")
        for d in offered:
            print(f"        → {d.entity_id} {d.entity_name} · {d.strength} {d.confidence}%")

        print("\nПроверка текста резюме — файлы приложены, текст НЕ извлечён:")
        for kw, kind, note in RESUME_TEXT_CASES:
            ent = Entity(org_id=org.id, type=EntityType.candidate, created_by=author_id,
                         status=EntityStatus.new, **kw)
            db.add(ent)
            await db.flush()
            # ОДИН И ТОТ ЖЕ текст резюме на всю группу (меняется только шапка с
            # именем) — так проверяется именно детектор копипаста.
            shared = _resume_blocks("Копия резюме", "Senior Recruiter", "Финтех-компания", "Москва")
            blocks = [(TITLE, ent.name)] + shared[1:]
            if kind == "pdf-shared":
                _attach_file(db, ent, org.id, _render_resume_pdf(blocks), f"resume_{ent.id}.pdf", "application/pdf")
            elif kind == "scan":
                _attach_file(db, ent, org.id, _scan_pdf(blocks), f"scan_{ent.id}.pdf", "application/pdf")
            elif kind == "docx":
                _attach_file(db, ent, org.id, _docx_bytes(blocks), f"resume_{ent.id}.docx",
                             "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
            print(f"  [{ent.id}] {ent.name:24} — {note}")

        await db.commit()

    await engine.dispose()
    print("\nГотово. Убрать демо-данные: .venv/bin/python seed_duplicates_demo.py --purge")


if __name__ == "__main__":
    asyncio.run(main())
