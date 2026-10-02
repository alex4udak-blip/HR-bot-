"""Резюме-картинка внутри Word/PDF (Диана Булатова, 02.10.2026).

Рекрутёр загрузил .docx, в котором всё резюме — одна вставленная картинка:
текста в файле ноль, и разбор падал с «Document appears to be empty or
unreadable». Та же картинка, загруженная как JPG, читалась нормально — у
картинок есть распознавание через Claude Vision. Теперь документ без текста
тоже идёт в распознавание: из .docx достаём вложенные картинки, у PDF рисуем
страницы. Платный запрос уходит ТОЛЬКО когда текста в файле нет.
"""
import io

import pytest
from PIL import Image

from api.services.documents import (
    DocumentParseResult,
    DocumentParser,
    OCR_MAX_IMAGE_BYTES,
    OCR_MAX_IMAGE_SIDE,
)


def _jpeg(width: int = 1200, height: int = 1700, noise: bool = True) -> bytes:
    """Картинка «как скан»: однотонная жмётся в десяток байт и отсеялась бы
    как логотип, поэтому рисуем шум — вес получается как у настоящего скана."""
    img = Image.new("RGB", (width, height), "white")
    if noise:
        px = img.load()
        for x in range(0, width, 2):
            for y in range(0, height, 2):
                px[x, y] = ((x * 7) % 256, (y * 13) % 256, (x * y) % 256)
    out = io.BytesIO()
    img.save(out, format="JPEG", quality=92)
    return out.getvalue()


def _docx(*, text: str = "", image: bytes | None = None) -> bytes:
    from docx import Document

    doc = Document()
    if text:
        doc.add_paragraph(text)
    if image is not None:
        doc.add_picture(io.BytesIO(image))
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _scanned_pdf(image: bytes, pages: int = 1) -> bytes:
    import pymupdf

    doc = pymupdf.open()
    for _ in range(pages):
        page = doc.new_page(width=595, height=842)
        page.insert_image(pymupdf.Rect(0, 0, 595, 842), stream=image)
    data = doc.tobytes()
    doc.close()
    return data


@pytest.fixture
def parser_with_fake_ocr():
    """Парсер, у которого Vision подменён: в тестах в Claude не ходим."""
    parser = DocumentParser()
    calls = []

    async def fake_ocr(file_bytes: bytes, filename: str) -> DocumentParseResult:
        calls.append({"filename": filename, "size": len(file_bytes)})
        return DocumentParseResult(
            content="ДИАНА БУЛАТОВА\nМаркетолог / таргетолог",
            status="parsed",
        )

    parser._ocr_with_vision = fake_ocr
    return parser, calls


@pytest.mark.asyncio
async def test_docx_with_only_image_is_recognized(parser_with_fake_ocr):
    parser, calls = parser_with_fake_ocr
    result = await parser.parse(_docx(image=_jpeg()), "Diana_Bulatova.docx")

    assert result.status == "parsed"
    assert "ДИАНА БУЛАТОВА" in result.content
    assert result.metadata["ocr_source"] == "docx"
    assert len(calls) == 1


@pytest.mark.asyncio
async def test_docx_with_text_does_not_pay_for_ocr(parser_with_fake_ocr):
    """Обычное резюме текстом — распознавание не трогаем (оно платное)."""
    parser, calls = parser_with_fake_ocr
    result = await parser.parse(
        _docx(text="Булатова Диана, маркетолог", image=_jpeg()), "resume.docx"
    )

    assert "Булатова Диана" in result.content
    assert calls == []


@pytest.mark.asyncio
async def test_docx_with_tiny_images_only_falls_back(parser_with_fake_ocr):
    """Пустой Word с логотипом-мелочью: распознавать нечего, OCR не зовём."""
    parser, calls = parser_with_fake_ocr
    logo = _jpeg(width=60, height=60, noise=False)
    result = await parser.parse(_docx(image=logo), "empty.docx")

    assert calls == []
    assert not result.content.strip()


@pytest.mark.asyncio
async def test_scanned_pdf_is_recognized(parser_with_fake_ocr):
    parser, calls = parser_with_fake_ocr
    result = await parser.parse(_scanned_pdf(_jpeg()), "scan.pdf")

    assert result.status == "parsed"
    assert "ДИАНА БУЛАТОВА" in result.content
    assert result.metadata["ocr_source"] == "pdf"
    assert len(calls) == 1


@pytest.mark.asyncio
async def test_multipage_scan_numbers_pages(parser_with_fake_ocr):
    parser, calls = parser_with_fake_ocr
    result = await parser.parse(_scanned_pdf(_jpeg(), pages=2), "scan2.pdf")

    assert len(calls) == 2
    assert "--- 1 ---" in result.content and "--- 2 ---" in result.content


@pytest.mark.asyncio
async def test_images_are_shrunk_before_vision(parser_with_fake_ocr):
    """Vision не берёт больше ~5 МБ и сам ужимает сторону — крупный скан
    уменьшаем сами, иначе платим за лишние пиксели или получаем ошибку."""
    parser, calls = parser_with_fake_ocr
    big = _jpeg(width=4000, height=5000)
    assert len(big) > OCR_MAX_IMAGE_BYTES  # исходник действительно тяжёлый

    await parser.parse(_docx(image=big), "big.docx")

    assert calls and calls[0]["size"] < OCR_MAX_IMAGE_BYTES


def test_shrink_keeps_small_image_untouched():
    parser = DocumentParser()
    small = _jpeg(width=800, height=1000)
    suffix, data = parser._shrink_for_vision(small, ".jpg")

    assert data == small and suffix == ".jpg"


def test_shrink_limits_long_side():
    parser = DocumentParser()
    _, data = parser._shrink_for_vision(_jpeg(width=4000, height=5000), ".jpg")

    assert max(Image.open(io.BytesIO(data)).size) <= OCR_MAX_IMAGE_SIDE
